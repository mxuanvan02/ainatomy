#!/usr/bin/env python3
"""
Tạo Phiếu khảo sát chi tiết học sinh sau khi sử dụng AInatomy
Bao gồm: thông tin cá nhân, đánh giá từng tầng, cảm nghĩ chi tiết, góp ý
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_border(cell, **kwargs):
    """Set cell borders"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        if edge in kwargs:
            element = OxmlElement(f'w:{edge}')
            for attr, val in kwargs[edge].items():
                element.set(qn(f'w:{attr}'), str(val))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def create_detailed_survey():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PHIẾU KHẢO SÁT HỌC SINH")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Sau khi trải nghiệm công cụ AInatomy")
    run.italic = True
    run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    # Instructions box
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    p = cell.paragraphs[0]
    run = p.add_run("📋 Hướng dẫn: ")
    run.bold = True
    run.font.size = Pt(11)
    run = p.add_run("Em hãy trả lời trung thực các câu hỏi dưới đây. Không có câu trả lời đúng hay sai, đây là ý kiến cá nhân của em. Thông tin của em sẽ được bảo mật và chỉ sử dụng cho mục đích nghiên cứu giáo dục.")
    run.font.size = Pt(11)
    
    doc.add_paragraph()
    
    # ===== PHẦN A: THÔNG TIN CÁ NHÂN =====
    p = doc.add_paragraph()
    run = p.add_run("PHẦN A: THÔNG TIN CÁ NHÂN")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    # Info table
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    
    info_fields = [
        ("Họ và tên:", "________________________________________"),
        ("Mã học sinh:", "__________________"),
        ("Lớp:", "__________"),
        ("Trường:", "________________________________________"),
        ("Ngày làm phiếu:", "____/____/________"),
    ]
    
    for i, (label, value) in enumerate(info_fields):
        cell = table.cell(i, 0)
        p = cell.paragraphs[0]
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(11)
        
        cell = table.cell(i, 1)
        p = cell.paragraphs[0]
        run = p.add_run(value)
        run.font.size = Pt(11)
    
    doc.add_paragraph()
    
    # ===== PHẦN B: ĐÁNH GIÁ TẦNG 1 - HUẤN LUYỆN AI =====
    p = doc.add_paragraph()
    run = p.add_run("PHẦN B: ĐÁNH GIÁ TẦNG 1 - HUẤN LUYỆN AI")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 1: Sau khi trải nghiệm Tầng 1 (tạo dữ liệu, huấn luyện mô hình), em hiểu như thế nào về cách AI học?")
    run.bold = True
    
    options = [
        "☐ Em chưa hiểu gì về cách AI học",
        "☐ Em hiểu một chút nhưng còn nhiều điều chưa rõ",
        "☐ Em hiểu khá rõ: AI học từ dữ liệu, càng nhiều dữ liệu càng tốt",
        "☐ Em hiểu rất rõ: AI học từ dữ liệu, nhưng chất lượng dữ liệu quan trọng hơn số lượng",
        "☐ Em hiểu sâu: AI có thể bị thiên kiến nếu dữ liệu huấn luyện không cân bằng"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 2: Khi em tạo bộ dữ liệu với 80% ảnh ban ngày và 20% ảnh ban đêm, điều gì xảy ra với mô hình?")
    run.bold = True
    
    options = [
        "☐ Mô hình hoạt động tốt với cả ảnh ban ngày và ban đêm",
        "☐ Mô hình tốt với ảnh ban ngày nhưng kém với ảnh ban đêm",
        "☐ Mô hình kém với cả hai loại ảnh",
        "☐ Em không nhớ rõ",
        "☐ Khác: _________________________________________________"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 3: Em có bất ngờ khi thấy mô hình dự đoán sai với ảnh ban đêm không?")
    run.bold = True
    
    options = [
        "☐ Rất bất ngờ, em nghĩ AI luôn đúng",
        "☐ Hơi bất ngờ, nhưng em hiểu lý do",
        "☐ Không bất ngờ, em đã đoán được điều này",
        "☐ Em không có cảm xúc gì đặc biệt"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 4: Em cảm thấy thế nào khi tự tay tạo dữ liệu và huấn luyện mô hình?")
    run.bold = True
    
    options = [
        "☐ Rất thú vị, em muốn thử thêm",
        "☐ Thú vị, nhưng hơi phức tạp",
        "☐ Bình thường, không có gì đặc biệt",
        "☐ Khó hiểu, em cần giải thích thêm",
        "☐ Không thích, em thấy nhàm chán"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    # ===== PHẦN C: ĐÁNH GIÁ TẦNG 2 - KIỂM CHỨNG AI =====
    p = doc.add_paragraph()
    run = p.add_run("PHẦN C: ĐÁNH GIÁ TẦNG 2 - KIỂM CHỨNG AI")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 5: Trong Tầng 2, em đã phát hiện được bao nhiêu câu trả lời sai của AI (trong 12 câu)?")
    run.bold = True
    
    options = [
        "☐ 0-3 câu",
        "☐ 4-6 câu",
        "☐ 7-9 câu",
        "☐ 10-12 câu"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 6: Em dùng cách nào để phát hiện câu trả lời sai của AI? (Có thể chọn nhiều)")
    run.bold = True
    
    options = [
        "☐ Đối chiếu với kiến thức em đã biết",
        "☐ Tìm kiếm trên Internet để kiểm tra",
        "☐ Kiểm tra logic nội tại của câu trả lời",
        "☐ Xem nguồn thông tin AI trích dẫn",
        "☐ Hỏi bạn bè hoặc giáo viên",
        "☐ Dựa vào cảm giác cá nhân",
        "☐ Khác: _________________________________________________"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 7: Sau khi phát hiện AI sai, em cảm thấy thế nào?")
    run.bold = True
    
    options = [
        "☐ Bất ngờ, em nghĩ AI luôn đúng",
        "☐ Thất vọng, em mất niềm tin vào AI",
        "☐ Bình thường, em hiểu AI có thể sai",
        "☐ Thú vị, em muốn tìm hiểu thêm về lỗi của AI",
        "☐ Lo lắng, em sợ sử dụng AI trong tương lai"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 8: Em có nghĩ kỹ năng kiểm chứng AI là quan trọng không?")
    run.bold = True
    
    options = [
        "☐ Không quan trọng, AI luôn đáng tin cậy",
        "☐ Ít quan trọng, chỉ cần khi làm bài tập",
        "☐ Quan trọng, cần thiết trong học tập",
        "☐ Rất quan trọng, cần thiết trong cả học tập và cuộc sống",
        "☐ Cực kỳ quan trọng, em sẽ áp dụng ngay từ bây giờ"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    # ===== PHẦN D: ĐÁNH GIÁ TẦNG 3 - BẢN ĐỒ NĂNG LỰC =====
    p = doc.add_paragraph()
    run = p.add_run("PHẦN D: ĐÁNH GIÁ TẦNG 3 - BẢN ĐỒ NĂNG LỰC")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 9: Khi xem Bản đồ năng lực của mình, em thấy:")
    run.bold = True
    
    options = [
        "☐ Rất hữu ích, em biết được điểm mạnh và điểm yếu của mình",
        "☐ Hữu ích, nhưng em không chắc về độ chính xác",
        "☐ Bình thường, em không quan tâm lắm",
        "☐ Khó hiểu, em không biết đọc bản đồ như thế nào",
        "☐ Không hữu ích, em thấy không cần thiết"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 10: Năng lực nào em thấy mình mạnh nhất?")
    run.bold = True
    
    options = [
        "☐ Hiểu quy trình AI",
        "☐ Phân biệt loại AI (phân loại vs tạo sinh)",
        "☐ Kiểm chứng kết quả AI",
        "☐ Sử dụng công cụ AI",
        "☐ Tư duy phản biện về AI"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 11: Năng lực nào em thấy mình cần cải thiện nhất?")
    run.bold = True
    
    options = [
        "☐ Hiểu quy trình AI",
        "☐ Phân biệt loại AI (phân loại vs tạo sinh)",
        "☐ Kiểm chứng kết quả AI",
        "☐ Sử dụng công cụ AI",
        "☐ Tư duy phản biện về AI"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    # ===== PHẦN E: CẢM NGHĨ VÀ TRẢI NGHIỆM =====
    p = doc.add_paragraph()
    run = p.add_run("PHẦN E: CẢM NGHĨ VÀ TRẢI NGHIỆM")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 12: Điều gì làm em ấn tượng nhất khi sử dụng AInatomy?")
    run.bold = True
    
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 13: Điều gì làm em bất ngờ nhất?")
    run.bold = True
    
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 14: Phần nào em thấy thú vị nhất? Tại sao?")
    run.bold = True
    
    options = [
        "☐ Tầng 1: Huấn luyện AI - vì _________________________________________",
        "☐ Tầng 2: Kiểm chứng AI - vì _________________________________________",
        "☐ Tầng 3: Bản đồ năng lực - vì _________________________________________",
        "☐ Tất cả đều thú vị như nhau"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 15: Phần nào em thấy khó hiểu nhất? Tại sao?")
    run.bold = True
    
    options = [
        "☐ Tầng 1: Huấn luyện AI - vì _________________________________________",
        "☐ Tầng 2: Kiểm chứng AI - vì _________________________________________",
        "☐ Tầng 3: Bản đồ năng lực - vì _________________________________________",
        "☐ Không có phần nào khó hiểu",
        "☐ Khác: _________________________________________________"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 16: Em gặp khó khăn gì khi sử dụng AInatomy? (Có thể chọn nhiều)")
    run.bold = True
    
    options = [
        "☐ Không gặp khó khăn gì",
        "☐ Khó hiểu về khái niệm thiên kiến dữ liệu",
        "☐ Khó hiểu về cách kiểm chứng AI",
        "☐ Khó sử dụng giao diện công cụ",
        "☐ Khó hiểu về bản đồ năng lực",
        "☐ Công cụ chạy chậm",
        "☐ Khác: _________________________________________________"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    # ===== PHẦN F: ỨNG DỤNG VÀ TƯƠNG LAI =====
    p = doc.add_paragraph()
    run = p.add_run("PHẦN F: ỨNG DỤNG VÀ TƯƠNG LAI")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 17: Sau bài học này, khi sử dụng ChatGPT hoặc các công cụ AI khác, em sẽ:")
    run.bold = True
    
    options = [
        "☐ Tin ngay vào câu trả lời của AI",
        "☐ Vẫn tin nhưng sẽ kiểm tra lại một số thông tin quan trọng",
        "☐ Luôn kiểm tra lại câu trả lời của AI trước khi sử dụng",
        "☐ Không sử dụng AI nữa vì sợ bị sai",
        "☐ Sử dụng AI nhưng với thái độ thận trọng và phê phán"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 18: Em có muốn học thêm về AI trong các bài học tiếp theo không?")
    run.bold = True
    
    options = [
        "☐ Không muốn",
        "☐ Có nhưng chỉ một chút",
        "☐ Rất muốn học thêm",
        "☐ Muốn học chuyên sâu về AI"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 19: Nếu được học thêm về AI, em muốn học chủ đề gì? (Có thể chọn nhiều)")
    run.bold = True
    
    options = [
        "☐ Cách AI học và ra quyết định",
        "☐ Ứng dụng AI trong thực tế",
        "☐ Đạo đức AI và trách nhiệm",
        "☐ Lập trình AI cơ bản",
        "☐ AI trong ngành nghề cụ thể (y tế, giáo dục, nghệ thuật...)",
        "☐ Tương lai của AI và ảnh hưởng đến xã hội",
        "☐ Khác: _________________________________________________"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 20: Em sẽ giới thiệu AInatomy cho bạn bè hoặc người thân không?")
    run.bold = True
    
    options = [
        "☐ Không, em thấy không hữu ích",
        "☐ Có thể, nếu họ hỏi",
        "☐ Chắc chắn có, em thấy rất hữu ích",
        "☐ Chắc chắn có, em sẽ chia sẻ ngay"
    ]
    for opt in options:
        doc.add_paragraph(opt)
    
    doc.add_paragraph()
    
    # ===== PHẦN G: GÓP Ý CẢI THIỆN =====
    p = doc.add_paragraph()
    run = p.add_run("PHẦN G: GÓP Ý CẢI THIỆN")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 21: Theo em, điều gì cần cải thiện trong công cụ AInatomy?")
    run.bold = True
    
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 22: Em có muốn chia sẻ thêm điều gì về trải nghiệm học tập này không?")
    run.bold = True
    
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    
    doc.add_paragraph()
    
    # ===== PHẦN H: TỰ ĐÁNH GIÁ TỔNG THỂ =====
    p = doc.add_paragraph()
    run = p.add_run("PHẦN H: TỰ ĐÁNH GIÁ TỔNG THỂ")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 23: Em tự đánh giá mức độ hiểu bài của mình (khoanh tròn):")
    run.bold = True
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Chưa hiểu    ←    1    2    3    4    5    →    Hiểu rất rõ")
    run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 24: Em tự đánh giá mức độ hứng thú của em với bài học (khoanh tròn):")
    run.bold = True
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Không hứng thú    ←    1    2    3    4    5    →    Rất hứng thú")
    run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("Câu 25: Em tự đánh giá mức độ tự tin của em khi sử dụng AI sau bài học (khoanh tròn):")
    run.bold = True
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Không tự tin    ←    1    2    3    4    5    →    Rất tự tin")
    run.font.size = Pt(12)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Thank you
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("🎉 Cảm ơn em đã hoàn thành phiếu khảo sát!")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0, 128, 0)
    
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Ý kiến của em rất quý giá để cải thiện công cụ và phương pháp dạy học.")
    run.italic = True
    run.font.size = Pt(11)
    
    # Save
    output_path = "/home/hitokiri/ieeai2026/soi-ai/deliverables/teaching-trial/Phieu_khao_sat_chi_tiet.docx"
    doc.save(output_path)
    print(f"✓ Đã tạo phiếu khảo sát chi tiết: {output_path}")

if __name__ == "__main__":
    create_detailed_survey()
