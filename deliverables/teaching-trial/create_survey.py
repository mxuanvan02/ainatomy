#!/usr/bin/env python3
"""
Tạo Phiếu tự đánh giá học sinh sau khi sử dụng AInatomy
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_student_survey():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PHIẾU TỰ ĐÁNH GIÁ CỦA HỌC SINH")
    run.bold = True
    run.font.size = Pt(14)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Sau khi sử dụng công cụ AInatomy")
    run.italic = True
    run.font.size = Pt(11)
    
    doc.add_paragraph()
    
    # Student info
    p = doc.add_paragraph()
    run = p.add_run("Họ và tên: ")
    run.bold = True
    p.add_run("________________________________________")
    
    p = doc.add_paragraph()
    run = p.add_run("Lớp: ")
    run.bold = True
    p.add_run("__________________")
    
    run = p.add_run("    Ngày: ")
    run.bold = True
    p.add_run("__________________")
    
    doc.add_paragraph()
    
    # Instructions
    p = doc.add_paragraph()
    run = p.add_run("Hướng dẫn: ")
    run.bold = True
    p.add_run("Em hãy trả lời các câu hỏi sau một cách trung thực. Không có câu trả lời đúng hay sai, đây là ý kiến cá nhân của em.")
    
    doc.add_paragraph()
    
    # Question 1
    p = doc.add_paragraph()
    run = p.add_run("Câu 1: Sau khi sử dụng AInatomy, em hiểu như thế nào về cách AI hoạt động?")
    run.bold = True
    
    p = doc.add_paragraph("☐ Em chưa hiểu gì về cách AI hoạt động")
    p = doc.add_paragraph("☐ Em hiểu một chút nhưng còn nhiều điều chưa rõ")
    p = doc.add_paragraph("☐ Em hiểu khá rõ về quy trình AI (thu thập dữ liệu, huấn luyện, kiểm tra)")
    p = doc.add_paragraph("☐ Em hiểu rất rõ và có thể giải thích cho người khác")
    
    doc.add_paragraph()
    
    # Question 2
    p = doc.add_paragraph()
    run = p.add_run("Câu 2: Điều gì làm em bất ngờ nhất khi sử dụng AInatomy?")
    run.bold = True
    
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    
    doc.add_paragraph()
    
    # Question 3
    p = doc.add_paragraph()
    run = p.add_run("Câu 3: Sau bài học này, khi sử dụng ChatGPT hoặc các công cụ AI khác, em sẽ:")
    run.bold = True
    
    p = doc.add_paragraph("☐ Tin ngay vào câu trả lời của AI")
    p = doc.add_paragraph("☐ Vẫn tin nhưng sẽ kiểm tra lại một số thông tin quan trọng")
    p = doc.add_paragraph("☐ Luôn kiểm tra lại câu trả lời của AI trước khi sử dụng")
    p = doc.add_paragraph("☐ Không sử dụng AI nữa vì sợ bị sai")
    
    doc.add_paragraph()
    
    # Question 4
    p = doc.add_paragraph()
    run = p.add_run("Câu 4: Em nghĩ kỹ năng kiểm chứng AI có quan trọng không?")
    run.bold = True
    
    p = doc.add_paragraph("☐ Không quan trọng, AI luôn đúng")
    p = doc.add_paragraph("☐ Ít quan trọng, chỉ cần khi làm bài tập")
    p = doc.add_paragraph("☐ Quan trọng, cần thiết trong học tập")
    p = doc.add_paragraph("☐ Rất quan trọng, cần thiết trong cả học tập và cuộc sống")
    
    doc.add_paragraph()
    
    # Question 5
    p = doc.add_paragraph()
    run = p.add_run("Câu 5: Phần nào của AInatomy em thấy thú vị nhất?")
    run.bold = True
    
    p = doc.add_paragraph("☐ Tầng 1: Huấn luyện AI (tạo dữ liệu, xem mô hình học)")
    p = doc.add_paragraph("☐ Tầng 2: Kiểm chứng AI (phát hiện lỗi AI)")
    p = doc.add_paragraph("☐ Tầng 3: Bản đồ năng lực (xem điểm mạnh/yếu của em)")
    p = doc.add_paragraph("☐ Tất cả đều thú vị như nhau")
    
    doc.add_paragraph()
    
    # Question 6
    p = doc.add_paragraph()
    run = p.add_run("Câu 6: Em gặp khó khăn gì khi sử dụng AInatomy?")
    run.bold = True
    
    p = doc.add_paragraph("☐ Không gặp khó khăn gì")
    p = doc.add_paragraph("☐ Khó hiểu về khái niệm thiên kiến dữ liệu")
    p = doc.add_paragraph("☐ Khó hiểu về cách kiểm chứng AI")
    p = doc.add_paragraph("☐ Khó sử dụng giao diện công cụ")
    p = doc.add_paragraph("☐ Khác: _________________________________________________")
    
    doc.add_paragraph()
    
    # Question 7
    p = doc.add_paragraph()
    run = p.add_run("Câu 7: Em có muốn học thêm về AI trong các bài học tiếp theo không?")
    run.bold = True
    
    p = doc.add_paragraph("☐ Không muốn")
    p = doc.add_paragraph("☐ Có nhưng chỉ một chút")
    p = doc.add_paragraph("☐ Rất muốn học thêm")
    
    p = doc.add_paragraph()
    p = doc.add_paragraph("Lý do: _________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    
    doc.add_paragraph()
    
    # Question 8
    p = doc.add_paragraph()
    run = p.add_run("Câu 8: Em tự đánh giá mức độ hiểu bài của mình (khoanh tròn):")
    run.bold = True
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Chưa hiểu    ←    1    2    3    4    5    →    Hiểu rất rõ")
    run.font.size = Pt(11)
    
    doc.add_paragraph()
    
    # Question 9
    p = doc.add_paragraph()
    run = p.add_run("Câu 9: Theo em, điều gì cần cải thiện trong công cụ AInatomy?")
    run.bold = True
    
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    
    doc.add_paragraph()
    
    # Question 10
    p = doc.add_paragraph()
    run = p.add_run("Câu 10: Em có muốn chia sẻ thêm điều gì về trải nghiệm học tập này không?")
    run.bold = True
    
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    p = doc.add_paragraph("_________________________________________________________________")
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Thank you
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Cảm ơn em đã hoàn thành phiếu đánh giá!")
    run.italic = True
    run.font.size = Pt(11)
    
    # Save
    output_path = "/home/hitokiri/ieeai2026/soi-ai/deliverables/teaching-trial/Phieu_tu_danh_gia_HS.docx"
    doc.save(output_path)
    print(f"✓ Đã tạo phiếu tự đánh giá: {output_path}")

if __name__ == "__main__":
    create_student_survey()
