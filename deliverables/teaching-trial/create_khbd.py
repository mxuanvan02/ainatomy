#!/usr/bin/env python3
"""
Tạo Kế hoạch bài dạy (KHBD) theo CV 5512 cho việc dạy thử AInatomy
Tích hợp vào Bài 12 Tin học 10: "Phần mềm thiết kế đồ họa"
"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_border(cell, **kwargs):
    """Set cell borders"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        if edge in kwargs:
            element = OxmlElement(f'w:{edge}')
            for attr, val in kwargs[edge].items():
                element.set(qn(f'w:{attr}'), str(val))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_header_table(doc):
    """Add header with school info and national motto"""
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Remove borders
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, 
                top={"sz": 0, "val": "none"},
                bottom={"sz": 0, "val": "none"},
                start={"sz": 0, "val": "none"},
                end={"sz": 0, "val": "none"})
    
    # Left column - School info
    cell = table.cell(0, 0)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TRƯỜNG ĐẠI HỌC SƯ PHẠM, ĐẠI HỌC HUẾ\nTRƯỜNG THPT THUẬN HÓA")
    run.bold = True
    run.font.size = Pt(11)
    
    cell = table.cell(1, 0)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TỔ: KHOA HỌC TỰ NHIÊN")
    run.font.size = Pt(11)
    
    # Right column - National motto
    cell = table.cell(0, 1)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    run.bold = True
    run.font.size = Pt(11)
    
    cell = table.cell(1, 1)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Độc lập - Tự do - Hạnh phúc")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 128)
    
    # Add line under motto
    cell = table.cell(2, 1)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("_________________________")
    run.font.size = Pt(11)
    
    doc.add_paragraph()

def add_title(doc):
    """Add lesson plan title"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("KẾ HOẠCH BÀI DẠY")
    run.bold = True
    run.font.size = Pt(14)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Tên bài dạy: ỨNG DỤNG TRÍ TUỆ NHÂN TẠO TRONG THỰC TIỄN")
    run.bold = True
    run.font.size = Pt(13)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("(Tích hợp công cụ AInatomy vào Bài 12: Phần mềm thiết kế đồ họa)")
    run.italic = True
    run.font.size = Pt(11)
    
    doc.add_paragraph()

def add_info_table(doc):
    """Add teacher and class info"""
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    
    # Row 1
    table.cell(0, 0).text = "Giáo viên"
    table.cell(0, 1).text = "Mai Xuân Văn"
    
    # Row 2
    table.cell(1, 0).text = "Lớp"
    table.cell(1, 1).text = "10..."
    
    # Row 3
    table.cell(2, 0).text = "Số học sinh"
    table.cell(2, 1).text = ".... học sinh"
    
    # Row 4
    table.cell(3, 0).text = "Thời lượng"
    table.cell(3, 1).text = "2 tiết (90 phút)"
    
    # Bold first column
    for row in table.rows:
        for paragraph in row.cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    doc.add_paragraph()

def add_objectives(doc):
    """Add learning objectives"""
    p = doc.add_paragraph()
    run = p.add_run("I. YÊU CẦU CẦN ĐẠT")
    run.bold = True
    run.font.size = Pt(12)
    
    objectives = [
        "Về kiến thức:",
        "• Học sinh trình bày được quy trình hoạt động của hệ thống AI (thu thập dữ liệu, huấn luyện, triển khai, đánh giá).",
        "• Học sinh phân biệt được AI phân loại và AI tạo sinh, nêu được ví dụ cụ thể.",
        "• Học sinh giải thích được khái niệm thiên kiến dữ liệu và ảnh hưởng của nó đến chất lượng AI.",
        "• Học sinh mô tả được ít nhất 3 cách kiểm chứng đầu ra AI (đối chiếu nguồn, kiểm tra logic, thử nghiệm).",
        "",
        "Về năng lực:",
        "• Năng lực tư duy: Học sinh phân tích được đầu vào/đầu ra của hệ thống AI trong tình huống cụ thể.",
        "• Năng lực giải quyết vấn đề: Học sinh đề xuất được cách khắc phục khi phát hiện AI đưa ra kết quả sai.",
        "• Năng lực số: Học sinh sử dụng được công cụ AInatomy để trải nghiệm quy trình AI và kiểm chứng đầu ra.",
        "",
        "Về phẩm chất:",
        "• Trách nhiệm: Học sinh nhận thức được việc sử dụng AI có đạo đức và trách nhiệm.",
        "• Trung thực: Học sinh biết kiểm tra tính chính xác của thông tin do AI cung cấp.",
    ]
    
    for obj in objectives:
        p = doc.add_paragraph(obj)
        if obj.startswith("Về"):
            p.runs[0].bold = True
    
    doc.add_paragraph()

def add_preparation(doc):
    """Add preparation requirements"""
    p = doc.add_paragraph()
    run = p.add_run("II. CHUẨN BỊ")
    run.bold = True
    run.font.size = Pt(12)
    
    p = doc.add_paragraph()
    run = p.add_run("1. Giáo viên:")
    run.bold = True
    
    items_gv = [
        "• Máy tính có cài đặt trình duyệt web (Chrome/Firefox/Edge).",
        "• File AInatomy (chạy offline, không cần internet).",
        "• Phiếu học tập (Phụ lục 1).",
        "• Phiếu tự đánh giá của học sinh (Phụ lục 2).",
        "• Bảng kiểm đánh giá năng lực (Phụ lục 3).",
    ]
    for item in items_gv:
        doc.add_paragraph(item)
    
    p = doc.add_paragraph()
    run = p.add_run("2. Học sinh:")
    run.bold = True
    
    items_hs = [
        "• Ôn lại kiến thức về đồ họa điểm ảnh và đồ họa vectơ (Bài 12).",
        "• Chuẩn bị tâm thế sẵn sàng trải nghiệm công cụ AI.",
    ]
    for item in items_hs:
        doc.add_paragraph(item)
    
    doc.add_paragraph()

def add_teaching_process(doc):
    """Add teaching activities"""
    p = doc.add_paragraph()
    run = p.add_run("III. TIẾN TRÌNH DẠY HỌC")
    run.bold = True
    run.font.size = Pt(12)
    
    # Activity 1: Warm-up
    p = doc.add_paragraph()
    run = p.add_run("Hoạt động 1: Khởi động (10 phút)")
    run.bold = True
    run.font.size = Pt(11)
    
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    
    # Header
    table.cell(0, 0).text = "Hoạt động của giáo viên"
    table.cell(0, 1).text = "Hoạt động của học sinh"
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    # Content
    table.cell(1, 0).text = """• Chiếu 2 hình ảnh: một logo thiết kế bằng tay và một logo do AI tạo sinh.
• Đặt câu hỏi: "Theo em, hình ảnh nào do con người tạo ra, hình ảnh nào do AI tạo? Làm sao em biết?"
• Dẫn dắt vào bài: "Hôm nay chúng ta sẽ tìm hiểu cách AI hoạt động và làm thế nào để kiểm chứng kết quả do AI tạo ra."
• Giới thiệu công cụ AInatomy: "Đây là phòng thí nghiệm AI mô phỏng, giúp em trải nghiệm quy trình AI từ đầu đến cuối."""
    
    table.cell(1, 1).text = """• Quan sát 2 hình ảnh.
• Thảo luận cặp đôi (2 phút).
• Đại diện 2-3 cặp trả lời.
• Lắng nghe, ghi nhận vấn đề."""
    
    doc.add_paragraph()
    
    # Activity 2: Knowledge formation
    p = doc.add_paragraph()
    run = p.add_run("Hoạt động 2: Hình thành kiến thức (35 phút)")
    run.bold = True
    run.font.size = Pt(11)
    
    p = doc.add_paragraph()
    run = p.add_run("2.1. Trải nghiệm AInatomy - Tầng 1: Huấn luyện AI (15 phút)")
    run.bold = True
    
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    
    table.cell(0, 0).text = "Hoạt động của giáo viên"
    table.cell(0, 1).text = "Hoạt động của học sinh"
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    table.cell(1, 0).text = """• Hướng dẫn học sinh mở AInatomy, vào Tầng 1.
• Yêu cầu: "Em hãy tạo bộ dữ liệu huấn luyện với 80% ảnh ban ngày, 20% ảnh ban đêm. Huấn luyện mô hình và quan sát độ chính xác."
• Hỏi: "Độ chính xác trên dữ liệu ban ngày là bao nhiêu? Trên dữ liệu ban đêm là bao nhiêu? Tại sao lại có sự chênh lệch?"
• Giải thích khái niệm thiên kiến dữ liệu.
• Yêu cầu học sinh lặp lại với bộ dữ liệu cân bằng 50-50. So sánh kết quả."""
    
    table.cell(1, 1).text = """• Mở AInatomy, vào Tầng 1.
• Thực hiện theo hướng dẫn:
  - Tạo bộ dữ liệu lệch (80% ngày, 20% đêm).
  - Huấn luyện mô hình.
  - Ghi nhận độ chính xác: ngày ~95%, đêm ~40%.
• Thảo luận: "Mô hình học chủ yếu từ ảnh ban ngày, nên kém với ảnh ban đêm."
• Lặp lại với bộ dữ liệu cân bằng.
• Ghi nhận: độ chính xác đêm tăng lên ~85%."""
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run("2.2. Trải nghiệm AInatomy - Tầng 2: Kiểm chứng AI (20 phút)")
    run.bold = True
    
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    
    table.cell(0, 0).text = "Hoạt động của giáo viên"
    table.cell(0, 1).text = "Hoạt động của học sinh"
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    table.cell(1, 0).text = """• Hướng dẫn học sinh vào Tầng 2: Đấu trường bắt lỗi AI.
• Yêu cầu: "Em hãy đọc 12 câu trả lời do AI tạo ra. Với mỗi câu, em hãy quyết định: câu này ĐÚNG hay SAI? Nếu sai, lỗi ở đâu?"
• Hỏi: "Em dùng cách nào để kiểm chứng? Em có tin ngay vào câu trả lời của AI không?"
• Tổng kết 3 cách kiểm chứng:
  1. Đối chiếu với nguồn đáng tin cậy.
  2. Kiểm tra logic nội tại.
  3. Thử nghiệm với tình huống khác."""
    
    table.cell(1, 1).text = """• Vào Tầng 2, đọc 12 câu trả lời.
• Phân tích từng câu:
  - Câu 1-4: Thông tin đúng, có căn cứ.
  - Câu 5-8: Thông tin sai, bịa đặt (hallucination).
  - Câu 9-12: Thông tin đúng nhưng thiếu nguồn.
• Ghi nhận: "AI có thể nói rất tự tin nhưng vẫn sai."
• Thảo luận nhóm: "Làm sao để biết câu nào đúng?"
• Trình bày 3 cách kiểm chứng."""
    
    doc.add_paragraph()
    
    # Activity 3: Practice
    p = doc.add_paragraph()
    run = p.add_run("Hoạt động 3: Luyện tập (25 phút)")
    run.bold = True
    run.font.size = Pt(11)
    
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    
    table.cell(0, 0).text = "Hoạt động của giáo viên"
    table.cell(0, 1).text = "Hoạt động của học sinh"
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    table.cell(1, 0).text = """• Chia lớp thành 4 nhóm.
• Giao nhiệm vụ: "Mỗi nhóm nhận một tình huống thực tế (Phụ lục 1). Em hãy:
  1. Xác định đây là AI phân loại hay AI tạo sinh.
  2. Phân tích đầu vào và đầu ra.
  3. Đề xuất cách kiểm chứng kết quả AI.
  4. Nếu AI sai, đề xuất cách khắc phục."
• Phát phiếu học tập cho từng nhóm.
• Theo dõi, hỗ trợ các nhóm."""
    
    table.cell(1, 1).text = """• Nhận nhiệm vụ theo nhóm.
• Thảo luận và điền vào phiếu học tập:
  - Nhóm 1: AI nhận diện khuôn mặt để điểm danh.
  - Nhóm 2: ChatGPT viết bài luận văn.
  - Nhóm 3: AI dự đoán điểm số học sinh.
  - Nhóm 4: AI tạo sinh hình ảnh quảng cáo.
• Phân tích đầu vào/đầu ra.
• Đề xuất cách kiểm chứng.
• Chuẩn bị trình bày (3 phút/nhóm)."""
    
    doc.add_paragraph()
    
    # Activity 4: Application
    p = doc.add_paragraph()
    run = p.add_run("Hoạt động 4: Vận dụng (20 phút)")
    run.bold = True
    run.font.size = Pt(11)
    
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    
    table.cell(0, 0).text = "Hoạt động của giáo viên"
    table.cell(0, 1).text = "Hoạt động của học sinh"
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    table.cell(1, 0).text = """• Yêu cầu các nhóm trình bày kết quả.
• Đặt câu hỏi phản biện: "Nhóm em đề xuất kiểm chứng bằng cách đối chiếu nguồn. Nhưng nếu nguồn cũng sai thì sao?"
• Hướng dẫn học sinh điền Phiếu tự đánh giá (Phụ lục 2).
• Thu thập dữ liệu từ AInatomy: "Mỗi em hãy xuất file JSON nhật ký học tập của mình và nộp cho giáo viên."
• Tổng kết bài học:
  - AI có thể sai, đặc biệt là AI tạo sinh.
  - Kiểm chứng là kỹ năng bắt buộc.
  - Trách nhiệm thuộc về người sử dụng AI."""
    
    table.cell(1, 1).text = """• Đại diện các nhóm trình bày (3 phút/nhóm).
• Các nhóm khác nhận xét, đặt câu hỏi.
• Trả lời câu hỏi phản biện.
• Điền Phiếu tự đánh giá:
  - Điều em hiểu nhất về AI.
  - Điều em thấy bất ngờ.
  - Em sẽ kiểm chứng AI như thế nào trong thực tế?
• Xuất file JSON từ AInatomy (Tầng 3 → Bản đồ năng lực → Xuất JSON).
• Nộp file JSON + Phiếu tự đánh giá."""
    
    doc.add_paragraph()

def add_assessment(doc):
    """Add assessment section"""
    p = doc.add_paragraph()
    run = p.add_run("IV. ĐÁNH GIÁ")
    run.bold = True
    run.font.size = Pt(12)
    
    p = doc.add_paragraph()
    run = p.add_run("1. Công cụ đánh giá:")
    run.bold = True
    
    items = [
        "• Phiếu học tập (đánh giá quá trình).",
        "• Phiếu tự đánh giá của học sinh (đánh giá nhận thức).",
        "• File JSON nhật ký từ AInatomy (đánh giá năng lực số).",
        "• Bảng kiểm năng lực (đánh giá tổng hợp).",
    ]
    for item in items:
        doc.add_paragraph(item)
    
    p = doc.add_paragraph()
    run = p.add_run("2. Tiêu chí đánh giá:")
    run.bold = True
    
    table = doc.add_table(rows=5, cols=3)
    table.style = 'Table Grid'
    
    # Header
    table.cell(0, 0).text = "Năng lực"
    table.cell(0, 1).text = "Tiêu chí"
    table.cell(0, 2).text = "Minh chứng"
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    # Content
    table.cell(1, 0).text = "Tư duy"
    table.cell(1, 1).text = "Phân tích được đầu vào/đầu ra của AI"
    table.cell(1, 2).text = "Phiếu học tập, câu trả lời trên lớp"
    
    table.cell(2, 0).text = "Giải quyết vấn đề"
    table.cell(2, 1).text = "Đề xuất được cách kiểm chứng AI"
    table.cell(2, 2).text = "Phiếu học tập, thảo luận nhóm"
    
    table.cell(3, 0).text = "Năng lực số"
    table.cell(3, 1).text = "Sử dụng được AInatomy, xuất được dữ liệu"
    table.cell(3, 2).text = "File JSON nhật ký"
    
    table.cell(4, 0).text = "Phẩm chất"
    table.cell(4, 1).text = "Trung thực khi kiểm chứng, trách nhiệm khi sử dụng AI"
    table.cell(4, 2).text = "Phiếu tự đánh giá, quan sát trên lớp"
    
    doc.add_paragraph()

def add_appendices(doc):
    """Add appendices"""
    p = doc.add_paragraph()
    run = p.add_run("V. PHỤ LỤC")
    run.bold = True
    run.font.size = Pt(12)
    
    doc.add_page_break()
    
    # Appendix 1: Student worksheet
    p = doc.add_paragraph()
    run = p.add_run("PHỤ LỤC 1: PHIẾU HỌC TẬP")
    run.bold = True
    run.font.size = Pt(12)
    
    p = doc.add_paragraph()
    run = p.add_run("Nhóm: ________________ Lớp: ________________")
    run.font.size = Pt(11)
    
    p = doc.add_paragraph()
    run = p.add_run("Tình huống: _____________________________________________")
    run.font.size = Pt(11)
    
    doc.add_paragraph()
    
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    
    table.cell(0, 0).text = "Câu hỏi"
    table.cell(0, 1).text = "Trả lời"
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    table.cell(1, 0).text = "1. Đây là AI phân loại hay AI tạo sinh? Tại sao?"
    table.cell(1, 1).text = ""
    
    table.cell(2, 0).text = "2. Đầu vào của AI là gì? Đầu ra là gì?"
    table.cell(2, 1).text = ""
    
    table.cell(3, 0).text = "3. Em đề xuất cách kiểm chứng kết quả AI như thế nào?"
    table.cell(3, 1).text = ""
    
    table.cell(4, 0).text = "4. Nếu AI sai, em đề xuất cách khắc phục như thế nào?"
    table.cell(4, 1).text = ""
    
    doc.add_page_break()
    
    # Appendix 2: Student self-assessment
    p = doc.add_paragraph()
    run = p.add_run("PHỤ LỤC 2: PHIẾU TỰ ĐÁNH GIÁ CỦA HỌC SINH")
    run.bold = True
    run.font.size = Pt(12)
    
    p = doc.add_paragraph()
    run = p.add_run("Họ và tên: ________________ Lớp: ________________")
    run.font.size = Pt(11)
    
    doc.add_paragraph()
    
    questions = [
        "1. Điều em hiểu nhất về AI sau bài học này là gì?",
        "_________________________________________________________________",
        "_________________________________________________________________",
        "",
        "2. Điều em thấy bất ngờ nhất về AI là gì?",
        "_________________________________________________________________",
        "_________________________________________________________________",
        "",
        "3. Sau bài học này, em sẽ kiểm chứng kết quả do AI tạo ra như thế nào trong thực tế?",
        "_________________________________________________________________",
        "_________________________________________________________________",
        "_________________________________________________________________",
        "",
        "4. Em tự đánh giá mức độ hiểu bài của mình (đánh dấu X):",
        "☐ Chưa hiểu    ☐ Hiểu cơ bản    ☐ Hiểu rõ    ☐ Hiểu sâu",
        "",
        "5. Em có muốn học thêm về AI không? Tại sao?",
        "_________________________________________________________________",
        "_________________________________________________________________",
    ]
    
    for q in questions:
        p = doc.add_paragraph(q)
        if q.startswith(("1.", "2.", "3.", "4.", "5.")):
            p.runs[0].bold = True

def main():
    """Main function to create the lesson plan"""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Add content
    add_header_table(doc)
    add_title(doc)
    add_info_table(doc)
    add_objectives(doc)
    add_preparation(doc)
    add_teaching_process(doc)
    add_assessment(doc)
    add_appendices(doc)
    
    # Save document
    output_path = "/home/hitokiri/ieeai2026/soi-ai/deliverables/teaching-trial/KHBD_AInatomy_Tin10.docx"
    doc.save(output_path)
    print(f"✓ Đã tạo kế hoạch bài dạy: {output_path}")

if __name__ == "__main__":
    main()
